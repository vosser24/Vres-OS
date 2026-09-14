"""Real parser/process tests plus explicitly scripted DB-boundary tests; no PostgreSQL claim."""
from __future__ import annotations

import io
import json
import math
import os
import subprocess
import sys
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest

from test_audit_regressions import ScriptedConnection
from vres_os import bootstrap, embeddings, hooks, ingestion, preferences, relations, repository, sources
from vres_os.artifacts import ArtifactService
from vres_os.config import ConfigStore, VresConfig
from vres_os.locking import local_lock, lock_is_held
from vres_os.processes import run_bounded
from vres_os.procedures import ProcedureService, fingerprint
from vres_os.refresh import validate_reviewed_delta
from vres_os.validation import artifact_manifest


def test_windows_data_path_has_no_second_vendor_directory(monkeypatch, tmp_path):
    from vres_os import paths
    monkeypatch.delenv('VRES_DATA_DIR', raising=False)
    def location(name, **kwargs):
        assert name == 'VresOS' and kwargs['appauthor'] is False
        return str(tmp_path / 'VresOS')
    monkeypatch.setattr(paths, 'user_data_dir', location)
    assert paths.data_dir() == tmp_path / 'VresOS'


@pytest.mark.parametrize('role,db', [(True, False), (False, True), (True, True)])
def test_provisioning_never_alters_existing_account_or_database(monkeypatch, role, db):
    conn = ScriptedConnection([('FROM pg_roles', role), ('FROM pg_database', db)])
    monkeypatch.setattr(bootstrap, '_ask', lambda *x: 'postgres')
    monkeypatch.setattr(bootstrap.getpass, 'getpass', lambda *x: 'ephemeral-test-value')
    monkeypatch.setattr(bootstrap, '_runtime_dsn', lambda **x: 'test-dsn')
    monkeypatch.setattr(bootstrap, '_driver', lambda: SimpleNamespace(connect=lambda *a, **kw: conn))
    monkeypatch.setitem(sys.modules, 'psycopg', SimpleNamespace(sql=object()))
    with pytest.raises(RuntimeError, match='never|will not'):
        bootstrap._provision_local_database(host='localhost', port=5432, database='already', runtime_user='already', sslmode='prefer')
    assert len(conn.calls) == 2 and not any('ALTER' in q or 'CREATE' in q for q, _ in conn.calls)


def test_failed_setup_restores_previous_config_and_secret(monkeypatch, tmp_path):
    store = ConfigStore(tmp_path / 'config.json')
    cfg = VresConfig(configured=True)
    store.save(cfg)
    vault = {cfg.database.password_key: 'previous-test-credential'}
    fake = SimpleNamespace(get=lambda k: vault.get(k), set=lambda k, v: vault.__setitem__(k, v), delete=lambda k: vault.pop(k, None))
    monkeypatch.setattr(bootstrap, 'ConfigStore', lambda: store)
    monkeypatch.setattr(bootstrap, 'SecretStore', lambda: fake)
    monkeypatch.setattr(bootstrap, '_ask', lambda prompt, default: default)
    monkeypatch.setattr(bootstrap, '_yes_no', lambda *a: False)
    monkeypatch.setattr(bootstrap.getpass, 'getpass', lambda *a: 'replacement-test-credential')
    monkeypatch.setattr(bootstrap, '_test_database', lambda **kw: None)
    monkeypatch.setattr(bootstrap, 'migrate', lambda: (_ for _ in ()).throw(RuntimeError('migration rejected')))
    with pytest.raises(SystemExit) as exc:
        bootstrap._interactive_setup()
    assert exc.value.code == 2
    assert store.load().configured is True
    assert vault[cfg.database.password_key] == 'previous-test-credential'
    assert 'credential' not in store.path.read_text()


def test_unknown_session_never_falls_through_to_other_window(monkeypatch):
    conn = ScriptedConnection([('FROM vres.sessions', None)])
    assert repository.Repository()._selected_task_id(conn, 7, 'unknown-session') is None
    assert len(conn.calls) == 1


def test_new_session_with_two_tasks_does_not_assume_last_focus():
    conn = ScriptedConnection([('SELECT id FROM vres.tasks', [{'id': 1}, {'id': 2}])])
    assert repository.Repository()._selected_task_id(conn, 7, None) is None


def test_open_session_retry_is_idempotent_and_does_not_select_task(monkeypatch):
    conn = ScriptedConnection([('pg_advisory_xact_lock', None), ('SELECT session_key', {'session_key': 'S-1'})])
    monkeypatch.setattr(repository, 'connect', lambda: conn)
    assert repository.Repository().open_session(7, 'native-1') == 'S-1'


def test_binding_unregistered_session_fails_before_changing_focus(monkeypatch):
    conn = ScriptedConnection([('SELECT id FROM vres.tasks', {'id': 2}), ('SELECT project_id,status', {'project_id': 7, 'status': 'active'}), ('UPDATE vres.sessions', None)])
    monkeypatch.setattr(repository, 'connect', lambda: conn)
    with pytest.raises(ValueError, match='unregistered'):
        repository.Repository().bind_session(7, 'native-missing', 'T-2')


def test_prompt_hook_delivers_current_session_even_without_task(monkeypatch, capsys):
    monkeypatch.setattr(hooks, '_input', lambda: {'session_id': 'AFTER-CLEAR', 'prompt': 'continue'})
    monkeypatch.setattr(hooks, 'ConfigStore', lambda: SimpleNamespace(load=lambda: VresConfig(configured=True)))
    monkeypatch.setattr(hooks, '_project_id', lambda *a: 7)
    opened = []
    monkeypatch.setattr(hooks, 'Repository', lambda: SimpleNamespace(open_session=lambda *a: opened.append(a), active_task=lambda *a: None))
    hooks.user_prompt()
    value = json.loads(capsys.readouterr().out)
    assert 'AFTER-CLEAR' in value['hookSpecificOutput']['additionalContext']
    assert opened == [(7, 'AFTER-CLEAR')]


def test_child_hook_cannot_write_main_session_state(monkeypatch, capsys):
    monkeypatch.setattr(hooks, '_input', lambda: {'agent_id': 'child', 'session_id': 'S', 'prompt': 'do work'})
    monkeypatch.setattr(hooks, 'Repository', lambda: pytest.fail('child accessed main repository'))
    hooks.user_prompt(); hooks.session_start(); hooks.post_compact()
    assert capsys.readouterr().out == ''


def test_local_lock_uses_kernel_and_releases_on_exception(monkeypatch, tmp_path):
    monkeypatch.setenv('VRES_DATA_DIR', str(tmp_path))
    with pytest.raises(RuntimeError):
        with local_lock('test') as acquired:
            assert acquired and lock_is_held('test')
            raise RuntimeError('crash simulation')
    assert not lock_is_held('test')


def test_local_lock_blocks_another_actual_process(monkeypatch, tmp_path):
    monkeypatch.setenv('VRES_DATA_DIR', str(tmp_path))
    monkeypatch.setenv('PYTHONPATH', str(Path(__file__).parents[1] / 'src'))
    script = "from vres_os.locking import local_lock\nwith local_lock('process-test') as a: print(a)"
    with local_lock('process-test') as acquired:
        assert acquired
        r = subprocess.run([sys.executable, '-c', script], capture_output=True, text=True, check=True, timeout=5)
        assert r.stdout.strip() == 'False'
    r = subprocess.run([sys.executable, '-c', script], capture_output=True, text=True, check=True, timeout=5)
    assert r.stdout.strip() == 'True'


def test_bounded_process_round_trips_greek_and_redacts():
    result = run_bounded([sys.executable, '-c', "import sys; print(sys.stdin.read()); print('password=secret123')"], prompt='Ελληνικά', timeout=5)
    assert result.ok and 'Ελληνικά' in result.output and 'secret123' not in result.output


def test_bounded_process_timeout_is_not_reported_as_success():
    result = run_bounded([sys.executable, '-c', 'import time; time.sleep(20)'], timeout=.1)
    assert not result.ok and result.returncode == 124


def test_bounded_process_limits_fast_output_and_retained_tail():
    result = run_bounded([sys.executable, '-c', "print('x'*100000)"], max_output_bytes=4096, max_result_chars=100, timeout=5)
    assert not result.ok and result.returncode == 125 and len(result.output) < 300


@pytest.mark.parametrize('vector', [[], [0, 0], [math.inf], [math.nan]])
def test_bad_vectors_cannot_enter_index(vector):
    with pytest.raises(ValueError):
        embeddings._validate_vector(vector)


def test_instruct_embedding_is_not_silently_wrong_adapter():
    with pytest.raises(embeddings.EmbeddingUnavailable):
        embeddings.embedding_text('intfloat/multilingual-e5-large-instruct', 'hello', query=True)


def test_embedding_load_failure_returns_claims_to_queue(monkeypatch):
    svc = embeddings.EmbeddingService()
    cfg = VresConfig(embeddings_enabled=True)
    monkeypatch.setattr(embeddings, 'ConfigStore', lambda: SimpleNamespace(load=lambda: cfg))
    monkeypatch.setattr(svc, 'queue_missing', lambda *a: 1)
    rows = [{'job_id': 1, 'chunk_id': 2, 'content': 'hello', 'claimed_attempt': 1}]
    monkeypatch.setattr(svc, '_claim', lambda *a: rows)
    monkeypatch.setattr(embeddings, '_load_model', lambda *a: (_ for _ in ()).throw(RuntimeError('offline')))
    failed = []
    monkeypatch.setattr(svc, '_fail', lambda r, e: failed.append((r, e)))
    with pytest.raises(embeddings.EmbeddingUnavailable):
        svc.run_pending()
    assert failed[0][0] == rows


def test_stale_embedding_worker_cannot_publish_over_newer_claim(monkeypatch):
    svc = embeddings.EmbeddingService(); cfg = VresConfig(embeddings_enabled=True)
    monkeypatch.setattr(embeddings, 'ConfigStore', lambda: SimpleNamespace(load=lambda: cfg))
    monkeypatch.setattr(svc, 'queue_missing', lambda *a: 1)
    monkeypatch.setattr(svc, '_claim', lambda *a: [{'job_id': 1, 'chunk_id': 2, 'content': 'hello', 'claimed_attempt': 1}])
    monkeypatch.setattr(embeddings, '_load_model', lambda *a: SimpleNamespace(encode=lambda *a, **kw: [[1.0, 0.0]]))
    conn = ScriptedConnection([('pg_extension', {'ok': False}), ('SELECT attempts,status', {'attempts': 2, 'status': 'running'})])
    monkeypatch.setattr(embeddings, '_connect', lambda: conn)
    result = svc.run_pending()
    assert result['processed'] == 0


def test_source_dedup_query_contains_scope_authority_and_version():
    conn = ScriptedConnection([('pg_advisory_xact_lock', None), ('IS NOT DISTINCT FROM', {'id': 3, 'source_key': 'SRC-existing'}), ('INSERT INTO vres.source_locations', None)])
    key, sid = sources.SourceService().register_in_conn(conn, source_type='legacy', title='x', content_hash='a'*64, project_id=4, authority_level='unknown', version='2', path_or_uri='file.txt')
    assert (key, sid) == ('SRC-existing', 3)
    assert conn.calls[1][1] == ('a'*64, 4, 'legacy', 'unknown', '2')
    assert 'project_id=excluded.project_id' not in conn.calls[2][0]


def test_source_uri_strips_credentials_without_breaking_ipv6():
    uri = sources._safe_uri('https://user:secret@[::1]:8443/docs?token=123&lang=el')
    assert 'user:' not in uri and 'secret' not in uri and '123' not in uri
    assert '[::1]:8443' in uri and 'lang=el' in uri


def test_unknown_graph_object_is_not_zero_impact(monkeypatch):
    conn = ScriptedConnection([('SELECT kind,key,project_id', [])])
    monkeypatch.setattr(relations, '_connect', lambda: conn)
    with pytest.raises(KeyError): relations.impact('unknown', project_id=1)


def test_known_graph_object_without_edges_is_zero_impact(monkeypatch):
    conn = ScriptedConnection([('SELECT kind,key,project_id', [{'kind': 'registry', 'key': 'K', 'project_id': 1}]), ('FROM vres.relations r JOIN', [])])
    monkeypatch.setattr(relations, '_connect', lambda: conn)
    assert relations.impact('K', project_id=1) == []


def test_ambiguous_graph_key_must_specify_kind(monkeypatch):
    conn = ScriptedConnection([('SELECT kind,key,project_id', [{'kind': 'registry', 'key': 'K'}, {'kind': 'knowledge', 'key': 'K'}])])
    monkeypatch.setattr(relations, '_connect', lambda: conn)
    with pytest.raises(ValueError, match='Ambiguous'): relations.impact('K', project_id=1)


def test_relations_keep_additive_provenance(monkeypatch):
    conn = ScriptedConnection([('vres.registry_objects', {'id': 1, 'project_id': 7}), ('vres.registry_objects', {'id': 2, 'project_id': 7}), ('INSERT INTO vres.relations', {'id': 9}), ('INSERT INTO vres.relation_evidence', None)])
    monkeypatch.setattr(relations, '_connect', lambda: conn)
    relations.relate('registry', 'A', 'affects', 'registry', 'B', provenance='second source')
    assert 'provenance=excluded' not in conn.calls[2][0]


def test_preference_cannot_remove_negation_from_user_quote(monkeypatch):
    conn = ScriptedConnection([('pg_advisory_xact_lock', None), ('SELECT e.id,e.payload', {'id': 1, 'project_id': 7, 'payload': {'text': 'Do not sort by Units'}})])
    monkeypatch.setattr(preferences, '_connect', lambda: conn)
    with pytest.raises(ValueError, match='quote'):
        preferences.PreferenceService().set('sort', 'sort by Units', task_key='T')


def test_nonexistent_artifact_is_not_registered(tmp_path):
    with pytest.raises(FileNotFoundError):
        ArtifactService().register(title='output', artifact_type='report', path=str(tmp_path/'missing'), project_id=1)


def test_baseline_fingerprint_includes_all_meaning_and_is_order_stable():
    assert fingerprint({'a': 1, 'b': 2}) == fingerprint({'b': 2, 'a': 1})
    assert fingerprint({'output': {'top': 'Units'}}) != fingerprint({'output': {'top': 'Net Sales'}})


def test_office_preflight_rejects_traversal_and_duplicates(tmp_path):
    p = tmp_path/'unsafe.docx'
    with zipfile.ZipFile(p, 'w') as z: z.writestr('../document.xml', 'bad')
    with pytest.raises(ingestion.UnsafeContainer): ingestion._ooxml_preflight(p)


def test_docx_parser_reads_real_unicode_document(tmp_path):
    from docx import Document
    p = tmp_path/'Διαδικασία.docx'; d = Document(); d.add_paragraph('Έγκριση προμηθευτή'); d.save(p)
    doc = ingestion.extract(p)
    assert 'Έγκριση προμηθευτή' in doc.text


def test_excel_parser_ignores_other_sheets_only_if_recipe_says_so_not_ingestion(tmp_path):
    # Ingestion catalogs every bounded sheet; a learned reporting recipe selects its own first sheet.
    from openpyxl import Workbook
    p = tmp_path/'test.xlsx'; w=Workbook(); w.active.append(['SKU','Πωλήσεις']); w.active.append(['000042',10]); w.create_sheet('B').append(['second-sheet']); w.save(p)
    doc = ingestion.extract(p)
    assert '000042' in doc.text and 'second-sheet' in doc.text


def test_pdf_without_text_requires_review_not_invented_knowledge(tmp_path):
    from pypdf import PdfWriter
    p=tmp_path/'scan.pdf'; w=PdfWriter(); w.add_blank_page(width=100,height=100); w.write(p)
    with pytest.raises(ingestion.IngestionInputError): ingestion.extract(p)


def test_isolated_parser_executes_real_worker(monkeypatch, tmp_path):
    from docx import Document
    monkeypatch.setenv('PYTHONPATH', str(Path(__file__).parents[1]/'src'))
    p=tmp_path/'real.docx'; d=Document(); d.add_paragraph('Isolated test knowledge'); d.save(p)
    doc=ingestion.extract_isolated(p)
    assert 'Isolated test knowledge' in doc.text


def test_csv_semicolon_dialect_is_not_forced_to_comma(tmp_path):
    p=tmp_path/'greek.csv';p.write_text('SKU;Πωλήσεις\n000042;10\n',encoding='utf-8')
    doc=ingestion.extract(p)
    assert doc.metadata['delimiter']==';' and '000042' in doc.text


def test_raw_utf8_block_boundary_has_no_inserted_newline(tmp_path):
    p=tmp_path/'t.txt';text='α'*70000;p.write_text(text,encoding='utf-8')
    assert ingestion.extract(p).text == text


def test_onboarding_ignores_secrets_generated_dirs_and_symlinks(tmp_path):
    from vres_os.onboarding import _files
    (tmp_path/'readme.txt').write_text('ok');(tmp_path/'.env').write_text('secret')
    (tmp_path/'NODE_MODULES').mkdir();(tmp_path/'NODE_MODULES'/'bad.txt').write_text('ignore')
    (tmp_path/'link').symlink_to(tmp_path/'readme.txt')
    assert [p.name for p in _files(tmp_path)] == ['readme.txt']


def test_worker_does_not_import_project_shadow_package(monkeypatch, tmp_path):
    from vres_os.processes import worker_command
    from docx import Document
    p=tmp_path/'ok.docx'; d=Document(); d.add_paragraph('safe worker'); d.save(p)
    (tmp_path/'vres_os').mkdir()
    (tmp_path/'vres_os'/'__init__.py').write_text("raise RuntimeError('PROJECT SHADOW EXECUTED')")
    monkeypatch.setenv('PYTHONPATH', str(tmp_path))
    result = run_bounded(worker_command('vres_os.parse_worker', str(p)), cwd=tmp_path, timeout=10, redact_output=False)
    assert result.ok and 'safe worker' in result.output and 'PROJECT SHADOW EXECUTED' not in result.output


def test_approval_retry_requires_identical_procedure_contract(monkeypatch):
    from vres_os import procedures
    values = dict(procedure_key='P',name='Top 10',description='Accepted daily report',task_family='reports',project_id=7,
                  input_contract={'files':2},method=['append'],invariants=['SKU is text'],validation_contract=['totals reconcile'],
                  output_contract={'count':10},approval_key='APP')
    digest=fingerprint({'name':values['name'],'description':values['description'],'family':values['task_family'],
                        'input':values['input_contract'],'method':values['method'],'invariants':values['invariants'],
                        'validation':values['validation_contract'],'output':values['output_contract'],'implementation':None})
    monkeypatch.setattr(procedures,'require_approval',lambda *a: 9)
    conn=ScriptedConnection([('pg_advisory_xact_lock',None),('SELECT v.version_no',{'version_no':1,'contract_fingerprint':digest})])
    monkeypatch.setattr(procedures,'_connect',lambda:conn)
    assert ProcedureService().accept_baseline(**values)==('P',1)
    conn2=ScriptedConnection([('pg_advisory_xact_lock',None),('SELECT v.version_no',{'version_no':1,'contract_fingerprint':digest})])
    monkeypatch.setattr(procedures,'_connect',lambda:conn2)
    values['method']=['join instead']
    with pytest.raises(ValueError,match='new approval'):
        ProcedureService().accept_baseline(**values)


def test_same_candidate_retry_does_not_create_another_version(monkeypatch):
    from vres_os import procedures
    svc=ProcedureService()
    base={'version_no':1,'input_contract':{'files':2},'method':['append'],'invariants':['SKU text'],
          'validation_contract':['totals'], 'output_contract':{'rows':10}}
    monkeypatch.setattr(svc,'get',lambda *a:{'version':base})
    metrics={'quality_score':1,'runtime_ms':10,'input_tokens':50,'output_tokens':50,'tokens':100,'runs':5}
    monkeypatch.setattr(svc,'preferred_metrics',lambda *a:metrics)
    conn=ScriptedConnection([('pg_advisory_xact_lock',None),('SELECT id,preferred_version',{'id':1,'preferred_version':1}),
                            ('SELECT candidate_version,decision,reason',{'candidate_version':2,'decision':'requires_user','reason':'held'})])
    monkeypatch.setattr(procedures,'_connect',lambda:conn)
    result=svc.evaluate_candidate(procedure_key='P',candidate={'method':['faster append']},metrics=metrics)
    assert result['candidate_version']==2 and result['replayed_request'] is True


def test_candidate_replay_refuses_stale_preferred_baseline(monkeypatch):
    from vres_os import procedures
    svc=ProcedureService()
    base={'version_no':1,'input_contract':{},'method':['a'],'invariants':['x'],'validation_contract':['v'],'output_contract':{}}
    monkeypatch.setattr(svc,'get',lambda *a:{'version':base})
    monkeypatch.setattr(svc,'preferred_metrics',lambda *a:{'quality_score':1,'runtime_ms':10,'input_tokens':50,'output_tokens':50,'tokens':100,'runs':5})
    conn=ScriptedConnection([('pg_advisory_xact_lock',None),('SELECT id,preferred_version',{'id':1,'preferred_version':2})])
    monkeypatch.setattr(procedures,'_connect',lambda:conn)
    with pytest.raises(ValueError,match='baseline changed'):
        svc.evaluate_candidate(procedure_key='P',candidate={},metrics={'quality_score':1,'runtime_ms':9,'input_tokens':50,'output_tokens':50})


def test_changed_reviewed_file_invalidates_completion(monkeypatch,tmp_path):
    from vres_os import validation
    state={'validation_status':'passed','objective':'test'}
    p=tmp_path/'out.txt';p.write_text('accepted')
    manifest=artifact_manifest(tmp_path,['out.txt']);p.write_text('changed')
    request={'id':1,'status':'passed','task_id':2,'objective':'test','state_digest':validation.state_digest(state),'artifact_manifest':manifest}
    conn=ScriptedConnection([('SELECT r.*,t.objective',request),('SELECT * FROM vres.task_state',state)])
    monkeypatch.setattr(validation,'_connect',lambda:conn)
    with pytest.raises(ValueError,match='files changed'):
        validation.ValidationService().assert_current('T',7,tmp_path)


def test_registry_change_cannot_move_identity_across_projects(monkeypatch):
    from vres_os import registry
    conn=ScriptedConnection([('pg_advisory_xact_lock',None),('SELECT object_type,project_id',{'object_type':'module','project_id':8})])
    monkeypatch.setattr(registry,'_connect',lambda:conn)
    with pytest.raises(ValueError,match='immutable'):
        registry.RegistryService().register(object_key='M',object_type='module',name='Module',project_id=7)


def test_empty_or_invalid_worker_output_budgets_are_rejected():
    with pytest.raises(ValueError,match='output limits'):
        run_bounded([sys.executable,'-c','print(1)'],max_output_bytes=-1)


@pytest.mark.parametrize("phase", ["validate", "validation", "review", "audit", " Validate "])
def test_all_validation_phase_aliases_keep_protected_model(phase):
    from vres_os.model_policy import ModelPolicyService
    result = ModelPolicyService().recommend(phase)
    assert result["model"] == "fable"
    assert result["effort"] == "high"
    assert result["protected"] is True
    assert result["fallback_allowed"] is False


def test_unknown_model_phase_cannot_guess_a_cheaper_model():
    from vres_os.model_policy import ModelPolicyService
    with pytest.raises(ValueError, match="Unknown model phase"):
        ModelPolicyService().recommend("validator2")


def test_mcp_exposes_full_procedure_not_only_search_summary():
    import ast
    root = Path(__file__).parents[1]
    tree = ast.parse((root / "src/vres_os/mcp_server.py").read_text())
    funcs = {n.name: n for n in tree.body if isinstance(n, ast.FunctionDef)}
    node = funcs["procedure_get"]
    assert any(isinstance(d, ast.Call) and isinstance(d.func, ast.Attribute)
               and d.func.attr == "tool" for d in node.decorator_list)
    text = ast.unparse(node)
    assert "_require_node" in text
    assert "ProcedureService().get" in text
